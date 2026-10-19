"""Reproduce supplied Excel arithmetic and examine neural-network energy.

New reconstruction, 2026-10-13. Input workbooks are read only.
Run: python energy_analysis.py --energy-xlsx ENERGY.xlsx --memory-xlsx MEMORY.xlsx
Requires: openpyxl, matplotlib. Standard-library arithmetic is used for the model.
No GPU power measurement or model inference is performed by this program.
"""
from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import json
import math
import operator
import re
from dataclasses import asdict, dataclass, replace
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter

ENERGY_PATH = Path(__file__).parent / 'inputs' / 'energy_source.xlsx'
MEMORY_PATH = Path(__file__).parent / 'inputs' / 'memory_source.xlsx'
SCALES = (1 / 16, 1 / 8, 1 / 4, 1 / 2, 1.0)
BATCHES = (1, 4, 8, 16)


class FormulaEvaluator:
    """Strict subset needed by these files: references, arithmetic, SUM(range).

    ast.parse reads syntax; it does not execute input Python. Only allowlisted
    AST nodes/operators are interpreted. Circular or unsupported formulas fail.
    This is deliberately not a general Excel calculation engine.
    """
    binary = {ast.Add: operator.add, ast.Sub: operator.sub,
              ast.Mult: operator.mul, ast.Div: operator.truediv,
              ast.Pow: operator.pow}

    def __init__(self, sheet):
        self.sheet = sheet
        self.memo = {}
        self.active = set()

    def cell(self, address):
        address = address.replace('$', '').upper()
        if address in self.memo:
            return self.memo[address]
        if address in self.active:
            raise ValueError(f'Circular reference: {address}')
        self.active.add(address)
        try:
            value = self.sheet[address].value
            if isinstance(value, str) and value.startswith('='):
                expression = value[1:].replace('$', '').replace('^', '**')
                expression = re.sub(
                    r'SUM\(([A-Z]+[0-9]+):([A-Z]+[0-9]+)\)',
                    lambda m: repr(self.sum_range(m.group(1), m.group(2))),
                    expression, flags=re.I)
                value = self.node(ast.parse(expression, mode='eval').body)
            self.memo[address] = value
            return value
        finally:
            self.active.remove(address)

    def number(self, value):
        if value is None:
            return 0  # Excel's blank-cell arithmetic rule, not missing input policy.
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return value
        raise TypeError(f'Non-numeric arithmetic input: {value!r}')

    def node(self, node):
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.Name) and re.fullmatch(r'[A-Za-z]+[0-9]+', node.id):
            return self.cell(node.id)
        if isinstance(node, ast.BinOp) and type(node.op) in self.binary:
            return self.binary[type(node.op)](self.number(self.node(node.left)),
                                             self.number(self.node(node.right)))
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            value = self.number(self.node(node.operand))
            return value if isinstance(node.op, ast.UAdd) else -value
        raise ValueError(f'Unsupported Excel formula node: {ast.dump(node)}')

    def sum_range(self, first, last):
        c1, r1, c2, r2 = range_boundaries(f'{first}:{last}')
        values = [self.cell(f'{get_column_letter(c)}{r}')
                  for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        return sum(v for v in values if isinstance(v, (int, float)))


@dataclass(frozen=True)
class Transformer:
    name: str
    vocabulary: int
    context: int
    width: int
    layers: int
    heads: int
    head_width: int
    batch: int = 1
    max_positions: int = 2048

    def validate(self):
        for key, value in asdict(self).items():
            if key != 'name' and (not isinstance(value, int) or value <= 0):
                raise ValueError(f'{key} must be a positive integer')
        if self.width != self.heads * self.head_width:
            raise ValueError('d_model != n_heads * d_head; standard MHA model inconsistent')
        if self.context > self.max_positions:
            raise ValueError('active context exceeds positional embedding capacity')


def model_from_sheet(sheet):
    return Transformer(str(sheet['B3'].value), int(sheet['C6'].value),
                       int(sheet['C7'].value), int(sheet['C8'].value),
                       int(sheet['C9'].value), int(sheet['C10'].value),
                       int(sheet['C11'].value), int(sheet['C5'].value), int(sheet['C7'].value))


def parameter_count(model):
    """Exact workbook parameter formula, including 5D FFN term.

    It is the supplied GPT-family analytical approximation. The paper labels
    D=4096,L=32 as Llama-7B; this is not a checkpoint-specific Llama parameterizer.
    """
    return (model.vocabulary * model.width + model.max_positions * model.width
            + model.layers * (4 * model.width * model.heads * model.head_width
                              + 8 * model.width ** 2 + 5 * model.width))


def full_prefix_flops(model):
    """Dense forward pass, workbook convention: multiplication+addition = 2 FLOPs."""
    b, s, d, layers = model.batch, model.context, model.width, model.layers
    return b * layers * (24 * s * d ** 2 + 4 * s ** 2 * d)


def decode_flops(model):
    """One new token, existing KV reused; all Q/K/V for new token still computed."""
    b, s, d, layers = model.batch, model.context, model.width, model.layers
    return b * layers * (24 * d ** 2 + 4 * s * d)


def decode_energy(model, weight_j_bit=4.5e-12, mac_j=1.5e-12,
                  kv_read_j_bit=4.5e-12, kv_write_j_bit=4.5e-12,
                  precision_bits=16):
    """Specified traffic model, not GPU wall-power measurement.

    Weight tensor read once per batch step. Existing S-1 KV entries read once,
    one new K/V entry written. Standard multi-head attention, no GQA/compression.
    Tied output LM head adds 2*B*D*V FLOPs without a second weight tensor.
    Activation/softmax/norm/leakage/IO omitted and explicitly reported.
    """
    model.validate()
    if precision_bits <= 0 or any(v < 0 for v in
        (weight_j_bit, mac_j, kv_read_j_bit, kv_write_j_bit)):
        raise ValueError('Invalid precision or negative energy coefficient')
    p = parameter_count(model)
    block_flops = decode_flops(model)
    head_flops = 2 * model.batch * model.width * model.vocabulary
    flops = block_flops + head_flops
    kv_factor = 2 * model.batch * model.width * model.layers * precision_bits
    components = {
        'weight_j': p * precision_bits * weight_j_bit,
        'mac_j': flops / 2 * mac_j,
        'kv_read_j': kv_factor * (model.context - 1) * kv_read_j_bit,
        'kv_write_j': kv_factor * kv_write_j_bit,
    }
    total = sum(components.values())
    return {**asdict(model), 'parameters': p, 'block_flops':block_flops,
            'lm_head_flops':head_flops, 'flops_per_batch_step': flops,
            'weight_bytes': p * precision_bits / 8,
            'kv_capacity_bytes': kv_factor * model.context / 8,
            **components, 'modeled_total_j_per_batch_step': total,
            'modeled_j_per_output_token': total / model.batch,
            'coverage': 'full weight read + block/LM-head MAC + KV read/write; no activation/norm/softmax/leakage/IO'}


def audit_workbook(path):
    formulas = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    results = []
    for sheet in formulas:
        evaluator = FormulaEvaluator(sheet)
        for row in sheet:
            for cell in row:
                if cell.data_type != 'f':
                    continue
                ref = cached[sheet.title][cell.coordinate].value
                record = {'file': path.name, 'sheet': sheet.title,
                          'cell': cell.coordinate, 'formula': cell.value,
                          'cached': ref}
                try:
                    actual = evaluator.cell(cell.coordinate)
                    numeric = isinstance(actual, (int, float)) and isinstance(ref, (int, float))
                    matched = math.isclose(actual, ref, rel_tol=1e-12, abs_tol=1e-12) if numeric else actual == ref
                    record.update(recalculated=actual, matches_cached=matched,
                                  absolute_error=abs(actual-ref) if numeric else None,
                                  status='match' if matched else 'mismatch')
                except (ValueError, TypeError, ZeroDivisionError, KeyError) as exc:
                    record.update(status='formula_error', error=str(exc), matches_cached=False)
                results.append(record)
    return formulas, results


def workbook_outputs(workbook):
    rows, issues, models = [], [], []
    for sheet in workbook:
        if sheet['B3'].value is None:
            continue
        model = model_from_sheet(sheet); models.append(model)
        ev = FormulaEvaluator(sheet)
        literal_no = ev.cell('F40'); literal_kv = ev.cell('Q43')
        read_j, write_j = ev.cell('C37'), ev.cell('C38')
        j_flop = ev.cell('D44')
        # Same spreadsheet cache equations, with layer count linked to C9.
        linked_kv = ev.cell('Q37') + ev.cell('Q40') + ev.cell('N29') * model.layers * j_flop
        rows.append({'sheet': sheet.title, 'model': model.name,
                     'parameters': parameter_count(model), 'batch': model.batch,
                     'context': model.context, 'layers': model.layers,
                     'no_cache_flops': ev.cell('F33'), 'cache_flops_literal': ev.cell('Q33'),
                     'j_per_flop_proxy': j_flop, 'memory_read_j_bit': read_j,
                     'memory_write_j_bit': write_j,
                     'literal_no_cache_j': literal_no, 'literal_cache_j': literal_kv,
                     'cache_with_linked_layers_j': linked_kv,
                     'weight_gib': ev.cell('C20') / 2**30,
                     'kv_gib': ev.cell('S15') / 2**30,
                     'no_cache_minus_cache_j': literal_no-literal_kv})
        if model.layers != 96:
            issues.append({'sheet': sheet.title, 'cells': 'N30,Q33',
                           'issue': f'Literal *96 differs from n_layers={model.layers}; linked-layer scenario reported separately.'})
        try:
            model.validate()
        except ValueError as exc:
            issues.append({'sheet': sheet.title, 'cells': 'C8,C10,C11', 'issue': str(exc)})
        if ev.cell('C31') != ev.cell('C30'):
            issues.append({'sheet': sheet.title, 'cells': 'C30,C31',
                           'issue': 'Hardcoded C31 differs from dynamic C30.'})
    return rows, issues, models


# Table 2 published normalized values. Preserve rounding rather than inventing
# hidden unrounded activation traffic or claiming exact absolute-J reproduction.
PAPER_BASE = {
    'ResNet18': {'weight': 0.65914, 'mac': 0.32393, 'activation': 0.01693, 'kv': 0.0},
    'Llama-7B': {'weight': 0.85374, 'mac': 0.01763, 'activation': 0.000447059, 'kv': 0.12819},
}
PAPER_TOTALS = {
    ('ResNet18','weight'): [[.38206,.42325,.50564,.67043,1], [1.40466,1.44585,1.52824,1.69303,2.02259], [2.76812,2.80931,2.8917,3.05649,3.38606], [5.49504,5.53623,5.61862,5.78341,6.11298]],
    ('ResNet18','mac'): [[.69631,.71656,.75705,.83803,1], [.80784,.88883,1.05079,1.37473,2.02259], [.95656,1.11852,1.44246,2.09033,3.38606], [1.25397,1.57791,2.22578,3.52151,6.11298]],
    ('Llama-7B','weight'): [[.19962,.25298,.3597,.57313,1],[.63841,.69176,.79848,1.01192,1.43879],[1.22345,1.27681,1.38353,1.59696,2.02383],[2.39355,2.44691,2.55363,2.76706,3.19393]],
    ('Llama-7B','mac'): [[.98348,.98458,.98678,.99119,1],[1.37269,1.37709,1.38591,1.40353,1.43879],[1.89164,1.90045,1.91808,1.95333,2.02383],[2.92954,2.94717,2.98242,3.05292,3.19393]],
    ('Llama-7B','kv'): [[.87982,.88783,.90386,.93591,1],[.95808,.99012,1.05422,1.18241,1.43879],[1.06242,1.12651,1.2547,1.51108,2.02383],[1.2711,1.39928,1.65566,2.16842,3.19393]],
}


def paper_sweep():
    rows = []
    for (model, target), expected in PAPER_TOTALS.items():
        base = PAPER_BASE[model]
        for bi,batch in enumerate(BATCHES):
            comp = {k:v*(1 if k=='weight' else batch) for k,v in base.items()}
            baseline = sum(comp.values())
            for si,scale in enumerate(SCALES):
                scaled = {k:v*(scale if k==target else 1) for k,v in comp.items()}
                total = sum(scaled.values()); reference = expected[bi][si]
                rows.append({'model':model, 'batch':batch, 'scaled_component':target,
                             'scale':scale, **scaled, 'normalized_total':total,
                             'paper_table2_total':reference,
                             'absolute_rounding_difference':abs(total-reference),
                             'reduction_percent':100*(1-total/baseline)})
    return rows


def memory_outputs(workbook):
    ev=FormulaEvaluator(workbook['Sheet1'])
    return [{'cell':cell, 'quantity':label, 'value':ev.cell(cell), 'unit':unit}
            for cell,label,unit in [
                ('E19','GPT parameter count','parameters'),
                ('E21','word + positional embedding','parameters'),
                ('E33','saved activation per layer, workbook FP16-equivalent count','FP16-equivalent elements'),
                ('E34','saved activation across layers, no parallel','FP16-equivalent elements'),
                ('E35','recomputation retained input per layer','elements'),
                ('E51','dense forward per layer','FLOPs')]]


def write_csv(path, records):
    if not records:
        return
    keys=list(dict.fromkeys(k for record in records for k in record))
    with path.open('w',encoding='utf-8-sig',newline='') as stream:
        writer=csv.DictWriter(stream,fieldnames=keys)
        writer.writeheader();writer.writerows(records)


def make_plots(out, literal, paper):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.size': 14, 'axes.labelsize': 14,
                         'axes.titlesize': 16, 'xtick.labelsize': 12,
                         'ytick.labelsize': 12, 'legend.fontsize': 12})
    colors={'weight':'#214E72','mac':'#078B8C','activation':'#B6C2CE','kv':'#E19937'}
    fig,axes=plt.subplots(1,2,figsize=(11,4.5))
    for ax,(model,base) in zip(axes,PAPER_BASE.items()):
        bottom=[0.0]*4
        for key in ('weight','mac','activation','kv'):
            values=[base[key]*(1 if key=='weight' else b) for b in BATCHES]
            ax.bar(range(4),values,bottom=bottom,label=key,color=colors[key])
            bottom=[a+b for a,b in zip(bottom,values)]
        ax.set_xticks(range(4),BATCHES);ax.set_xlabel('Batch size')
        ax.set_title(model);ax.set_ylabel('Energy / batch-1 baseline (a.u.)')
        ax.spines[['top','right']].set_visible(False)
    axes[1].legend(frameon=False);fig.tight_layout();fig.savefig(out/'paper_breakdown.png',dpi=180);plt.close(fig)
    fig,axes=plt.subplots(1,2,figsize=(11,4.5))
    for ax,model in zip(axes,PAPER_BASE):
        targets=['weight','mac']+(['kv'] if model=='Llama-7B' else [])
        width=.22
        for ti,target in enumerate(targets):
            vals=[next(r['reduction_percent'] for r in paper if r['model']==model and r['batch']==b and r['scaled_component']==target and r['scale']==1/16) for b in BATCHES]
            x=[i+(ti-(len(targets)-1)/2)*width for i in range(4)]
            ax.bar(x,vals,width=width,label=target,color=colors[target])
        ax.set_xticks(range(4),BATCHES);ax.set_ylim(0,100);ax.set_title(model)
        ax.set_xlabel('Batch size');ax.set_ylabel('Total energy reduction (%)')
        ax.spines[['top','right']].set_visible(False)
    axes[1].legend(frameon=False);fig.tight_layout();fig.savefig(out/'paper_reductions.png',dpi=180);plt.close(fig)
    fig,ax=plt.subplots(figsize=(10.5,4.5))
    x=range(len(literal));w=.25
    for shift,key,label,col in [(-w,'literal_no_cache_j','Excel: no cache','#214E72'),(0,'literal_cache_j','Excel: cache, literal *96','#D0763B'),(w,'cache_with_linked_layers_j','Same formula, linked layer count','#078B8C')]:
        ax.bar([v+shift for v in x],[r[key] for r in literal],width=w,label=label,color=col)
    ax.set_xticks(list(x),[r['model'] for r in literal]);ax.set_yscale('log');ax.set_ylabel('Workbook energy (J, log scale)')
    ax.legend(frameon=False,fontsize=11);ax.spines[['top','right']].set_visible(False)
    fig.tight_layout();fig.savefig(out/'workbook_energy.png',dpi=180);plt.close(fig)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--energy-xlsx',type=Path,default=ENERGY_PATH)
    parser.add_argument('--memory-xlsx',type=Path,default=MEMORY_PATH)
    parser.add_argument('--out',type=Path,default=Path(__file__).parent/'results')
    args=parser.parse_args();args.out.mkdir(parents=True,exist_ok=True)
    energy,audit1=audit_workbook(args.energy_xlsx)
    memory,audit2=audit_workbook(args.memory_xlsx)
    audit=audit1+audit2
    literal,issues,models=workbook_outputs(energy)
    paper=paper_sweep();decode=[]
    for model in models:
        for batch in BATCHES:
            try:
                decode.append(decode_energy(replace(model,batch=batch)))
            except ValueError:
                pass # Explicit issue recorded for invalid source dimensions.
    memory_rows=memory_outputs(memory)
    for name,rows in [('formula_audit',audit),('workbook_energy',literal),
                      ('paper_table2_reproduction',paper),('decode_scenarios',decode),
                      ('memory_sizing',memory_rows),('workbook_issues',issues)]:
        write_csv(args.out/(name+'.csv'),rows)
    max_diff=max(r['absolute_rounding_difference'] for r in paper)
    summary={'formula_cells':len(audit),'matches':sum(r['status']=='match' for r in audit),
             'formula_errors':[r for r in audit if r['status']!='match'],
             'workbook_models':len(literal),'paper_table2_cases':len(paper),
             'max_table2_rounding_difference':max_diff,'table2_rounding_tolerance':1e-4,
             'table2_rounding_check':max_diff<=1e-4,'consistent_decode_cases':len(decode),
             'workbook_issues':issues,
             'inputs':{str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in (args.energy_xlsx,args.memory_xlsx)},
             'units':{'workbook':'J; ratios /1024^3 are GiB, not decimal GB',
                      'paper_table2':'normalized to each model batch-1 baseline',
                      'decode':'modeled J per batch step; excludes activation/leakage/IO'},
             'interpretation':['Matching cached Excel values validates arithmetic, not model assumptions.',
                 'Excel cache mode removes K/V projections for all context tokens; not a standard one-token decode model.',
                 '250W/312TFLOP/s is a nameplate ratio, not measured per-operation energy.',
                 'KV capacity is linear in context length for fixed standard-MHA dimensions.',
                 'Paper Table 2 reproduction uses published rounded normalized components; it is not absolute-J reconstruction.']}
    (args.out/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    make_plots(args.out,literal,paper)
    print(json.dumps({k:v for k,v in summary.items() if k not in ('inputs','interpretation')},ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
