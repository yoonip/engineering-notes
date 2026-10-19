"""Excel 보조 단계의 오류가 CSV 결과를 훼손하지 않는지 검증한다."""
import csv
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from workbook_export import export_strict_workbook, optional_export


class WorkbookExportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.folder = Path(self.directory.name)
        self.summary = {"data": "unit-test fixture", "metrics": {"top_review_count": 1}, "training": {"selected_epoch": 2}}
        tables = {
            "model_comparison.csv": [["model", "f1"], ["fixture", 0.5]],
            "strict_causal_gpt_training_history.csv": [["epoch", "train_nll", "validation_nll"], [1, 0.2, 0.3], [2, 0.1, 0.2]],
            "normal_calibration_scores.csv": [["MaterialID", "score"], [3, 0.2]],
            "official_holdout_audit_scores.csv": [["MaterialID", "label"], [4, 0]],
            "saved_checkpoint_label_free_holdout_queue.csv": [["MaterialID", "review_rank", "anomaly_score"], [4, 1, 0.8], [5, 2, 0.5]],
        }
        for filename, rows in tables.items():
            with (self.folder / filename).open("w", encoding="utf-8", newline="") as handle:
                csv.writer(handle).writerows(rows)
        self.original_csvs = {path.name: path.read_bytes() for path in self.folder.glob("*.csv")}

    def tearDown(self):
        self.directory.cleanup()

    def assert_csvs_preserved(self):
        self.assertEqual(self.original_csvs, {path.name: path.read_bytes() for path in self.folder.glob("*.csv")})

    def test_without_node_writes_seven_sheets_and_blank_yield_formula(self):
        from openpyxl import load_workbook
        with patch.dict(os.environ, {"PATH": ""}), patch.object(subprocess, "run", side_effect=AssertionError("Node must not be invoked")):
            path, status = optional_export(lambda: export_strict_workbook(self.folder, self.summary))
        self.assertEqual(status["status"], "created")
        workbook = load_workbook(path)
        self.assertEqual(set(workbook.sheetnames), {"Summary", "Comparison", "PE follow-up", "Training", "Calibration", "Holdout audit", "Read me"})
        followup = workbook["PE follow-up"]
        self.assertEqual(followup.max_row, 2)
        self.assertEqual(followup["A2"].value, 4)
        self.assertIsNone(followup["J2"].value)
        self.assertIsNone(followup["K2"].value)
        self.assertIn('J2=""', followup["L2"].value)
        self.assertIn('J2>K2', followup["L2"].value)
        workbook.close()
        self.assert_csvs_preserved()

    def test_missing_openpyxl_retains_csv_results(self):
        with patch.dict(sys.modules, {"openpyxl": None}):
            path, status = optional_export(lambda: export_strict_workbook(self.folder, self.summary))
        self.assertIsNone(path)
        self.assertEqual(status["status"], "csv_only")
        self.assertIn("ModuleNotFoundError", status["reason"])
        self.assert_csvs_preserved()

    def test_unwritable_export_retains_csv_and_reports_failure(self):
        def denied():
            raise PermissionError("unit-test read-only destination")
        path, status = optional_export(denied)
        self.assertIsNone(path)
        self.assertEqual(status["status"], "csv_only")
        self.assertIn("PermissionError", status["reason"])
        self.assert_csvs_preserved()

    def test_missing_table_is_csv_only_not_false_success(self):
        (self.folder / "model_comparison.csv").unlink()
        path, status = optional_export(lambda: export_strict_workbook(self.folder, self.summary))
        self.assertIsNone(path)
        self.assertEqual(status["status"], "csv_only")
        self.assertFalse((self.folder / "strict_causal_gpt_experiment.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
