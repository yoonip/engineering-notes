"""CSV 결과를 공개 Python 라이브러리로 XLSX에 정리한다."""
from __future__ import annotations

import csv
import json
import math
from pathlib import Path
import tempfile
from typing import Callable


def optional_export(writer: Callable[[], Path]) -> tuple[Path | None, dict]:
    """Excel 단계가 실패해도 이미 저장한 모델·CSV 결과는 유지한다."""
    try:
        path = writer()
    except (ImportError, OSError, ValueError, KeyError) as error:
        return None, {"status": "csv_only", "reason": f"{type(error).__name__}: {error}"}
    return path, {"status": "created", "path": str(path)}


def _csv_rows(path: Path) -> list[list]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def _cell_value(value):
    if not isinstance(value, str):
        return value
    try:
        numeric = float(value)
        if math.isfinite(numeric):
            return int(numeric) if numeric.is_integer() else numeric
    except ValueError:
        pass
    return value


def export_strict_workbook(result_dir: Path, summary: dict) -> Path:
    """Transformer 결과의 지표·학습 이력·검토 큐를 7개 시트에 담는다."""
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill

    result_dir = Path(result_dir)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Summary"
    overview.append(["항목", "값"])
    for section in ("data", "task", "input", "unit_definition", "causality_scope", "research_evaluation_scope"):
        if section in summary:
            overview.append([section, summary[section]])
    for group in ("training", "metrics"):
        for key, value in summary.get(group, {}).items():
            overview.append([f"{group}.{key}", value])

    table_specs = [
        ("Comparison", "model_comparison.csv"),
        ("Training", "strict_causal_gpt_training_history.csv"),
        ("Calibration", "normal_calibration_scores.csv"),
        ("Holdout audit", "official_holdout_audit_scores.csv"),
    ]
    for sheet_name, filename in table_specs:
        sheet = workbook.create_sheet(sheet_name)
        for values in _csv_rows(result_dir / filename):
            sheet.append([_cell_value(value) for value in values])

    followup = workbook.create_sheet("PE follow-up", 2)
    followup.append(["MaterialID", "review_rank", "anomaly_score", "physical_sensor", "tool/chamber", "recipe", "WAT/CP/ATE evidence", "hypothesis", "DOE_plan", "good_dies", "tested_dies", "yield", "next_review"])
    queue_path = result_dir / "saved_checkpoint_label_free_holdout_queue.csv"
    with queue_path.open(encoding="utf-8-sig", newline="") as handle:
        queue = list(csv.DictReader(handle))
    capacity = int(summary.get("metrics", {}).get("top_review_count", max(1, math.ceil(len(queue) * 0.10))))
    queue.sort(key=lambda row: float(row["review_rank"]))
    for row in queue[:capacity]:
        followup.append([_cell_value(row["MaterialID"]), _cell_value(row["review_rank"]), _cell_value(row["anomaly_score"])])
        index = followup.max_row
        followup.cell(index, 12, f'=IF(OR(J{index}="",K{index}="",K{index}<=0,J{index}<0,J{index}>K{index}),"",J{index}/K{index})')
        followup.cell(index, 12).number_format = "0.00%"
        followup.cell(index, 13, f'=IF(G{index}="","검사 증거 확인",IF(I{index}="","DOE 계획 정리","조치 결과 검토"))')

    notes = workbook.create_sheet("Read me")
    notes.append(["항목", "내용"])
    notes.append(["Data", "STMicroelectronics ST-AWFD D2; https://github.com/STMicroelectronics/ST-AWFD"])
    notes.append(["License", "CC BY-NC-SA 4.0"])
    notes.append(["Scope", "Completed-step anomaly triage; no measured die-yield intervention outcome."])
    for item in summary.get("limitations", []):
        notes.append(["Evaluation", item])

    chart = LineChart()
    chart.title = "Normal-only Gaussian NLL"
    chart.x_axis.title = "Epoch"
    chart.y_axis.title = "Gaussian NLL"
    training = workbook["Training"]
    if training.max_row > 1 and training.max_column >= 3:
        chart.add_data(Reference(training, min_col=2, max_col=3, min_row=1, max_row=training.max_row), titles_from_data=True)
        chart.set_categories(Reference(training, min_col=1, min_row=2, max_row=training.max_row))
        training.add_chart(chart, "E2")

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="255C85")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(52, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
        # Source text beginning with '=' remains text; only the two explicit follow-up formulas execute.
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=") and not (sheet is followup and cell.column in (12, 13)):
                    cell.data_type = "s"

    destination = result_dir / "strict_causal_gpt_experiment.xlsx"
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=result_dir, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        temporary.replace(destination)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()
    return destination
